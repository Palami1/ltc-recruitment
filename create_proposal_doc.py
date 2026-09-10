import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Set margins to 2.54 cm (1 inch)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def set_font(run, name='Saysettha OT', size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    # Ensure East Asian / Complex script font is also set
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:cs'), name)
    rFonts.set(qn('w:eastAsia'), name)
    rPr.append(rFonts)

# Header
p_country = doc.add_paragraph()
p_country.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_country.paragraph_format.space_after = Pt(2)
p_country.paragraph_format.line_spacing = 1.15
r1 = p_country.add_run("ສາທາລະນະລັດ ປະຊາທິປະໄຕ ປະຊາຊົນລາວ\n")
set_font(r1, bold=True, size=13)
r2 = p_country.add_run("ສັນຕິພາບ ເອກະລາດ ປະຊາທິປະໄຕ ເອກະພາບ ວັດທະນະຖາວອນ\n---------------------")
set_font(r2, bold=True, size=13)

# Company & Department info in two columns table without border
table_header = doc.add_table(rows=2, cols=2)
table_header.alignment = WD_TABLE_ALIGNMENT.CENTER
table_header.autofit = False

# Left Column
cell_l1 = table_header.cell(0, 0)
cell_l1.width = Inches(3.2)
p_c1 = cell_l1.paragraphs[0]
p_c1.paragraph_format.space_after = Pt(2)
r_c1 = p_c1.add_run("ບໍລິສັດ ລາວ ໂທລະຄົມມະນາຄົມ ມະຫາຊົນ")
set_font(r_c1, bold=True, size=11)

cell_l2 = table_header.cell(1, 0)
cell_l2.width = Inches(3.2)
p_c2 = cell_l2.paragraphs[0]
p_c2.paragraph_format.space_after = Pt(2)
r_c2 = p_c2.add_run("ພະແນກຈັດຕັ້ງ")
set_font(r_c2, bold=True, size=11)

# Right Column
cell_r1 = table_header.cell(0, 1)
cell_r1.width = Inches(3.2)
p_r1 = cell_r1.paragraphs[0]
p_r1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_r1.paragraph_format.space_after = Pt(2)
r_num = p_r1.add_run("ເລກທີ: ......./ລລທ.ຈຕ")
set_font(r_num, size=11)

cell_r2 = table_header.cell(1, 1)
cell_r2.width = Inches(3.2)
p_r2 = cell_r2.paragraphs[0]
p_r2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_r2.paragraph_format.space_after = Pt(2)
r_date = p_r2.add_run("ນະຄອນຫຼວງວຽງຈັນ, ວັນທີ: .......................")
set_font(r_date, size=11)

# Title
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(14)
p_title.paragraph_format.space_after = Pt(12)
r_title = p_title.add_run("ໃບສະເໜີ")
set_font(r_title, bold=True, size=15)

# To and Subject
p_to = doc.add_paragraph()
p_to.paragraph_format.space_after = Pt(4)
p_to.paragraph_format.left_indent = Inches(0.5)
r_to1 = p_to.add_run("ຮຽນ: ")
set_font(r_to1, bold=True, size=12)
r_to2 = p_to.add_run("ທ່ານ ຮັກສາການຫົວໜ້າພະແນກໄອທີ ທີ່ນັບຖືຢ່າງສູງ.")
set_font(r_to2, size=12)

p_subj = doc.add_paragraph()
p_subj.paragraph_format.space_after = Pt(10)
p_subj.paragraph_format.left_indent = Inches(0.5)
r_s1 = p_subj.add_run("ເລື່ອງ: ")
set_font(r_s1, bold=True, size=12)
r_s2 = p_subj.add_run("ຂໍນຳໃຊ້ Server, ພື້ນທີ່ຈັດເກັບຂໍ້ມູນ ແລະ API ສຳລັບລະບົບຮັບສະໝັກພະນັກງານແບບອອນລາຍ (LTC-Recruitment).")
set_font(r_s2, bold=True, size=12)

# References
refs = [
    "ອີງຕາມ ຂໍ້ຕົກລົງ ຂອງ ຜູ້ອຳນວຍການໃຫຍ່ ລລທ ວ່າດ້ວຍການກຳນົດພາລະບົດບາດ, ສິດ ແລະ ໜ້າທີ່ການເຄື່ອນໄຫວວຽກງານ ຂອງ ພະແນກຈັດຕັ້ງ ສະບັບເລກທີ 1983/ລລທ, ລົງວັນທີ 09 ມິຖຸນາ 2022.",
    "ອີງຕາມ ໜັງສືສະເໜີຂໍພະນັກງານມາຊ່ວຍພັດທະນາເວັບໄຊ ຮັບສະໝັກວຽກພະນັກງານແບບອອນລາຍ ສະບັບເລກທີ 0141/ຈຕ.e, ລົງວັນທີ 20 ສິງຫາ 2026.",
    "ອີງຕາມ ບົດບັນທຶກກອງປະຊຸມປຶກສາຫາລືກ່ຽວກັບລະບົບສະໝັກວຽກພະນັກງານແບບອອນລາຍ (LTC-Recruitment) ລົງວັນທີ 08 ກັນຍາ 2026."
]

for ref in refs:
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.space_after = Pt(3)
    p_ref.paragraph_format.line_spacing = 1.15
    r = p_ref.add_run("- " + ref)
    set_font(r, size=11.5)

# Body Paragraph
p_body = doc.add_paragraph()
p_body.paragraph_format.space_before = Pt(8)
p_body.paragraph_format.space_after = Pt(8)
p_body.paragraph_format.line_spacing = 1.2
p_body.paragraph_format.first_line_indent = Inches(0.5)
r_b1 = p_body.add_run("ພະແນກຈັດຕັ້ງ ຈຶ່ງຂໍຮຽນສະເໜີມາຍັງ ທ່ານ ຮັກສາການຫົວໜ້າພະແນກໄອທີ ເພື່ອຂໍອະນຸມັດ ແລະ ອຳນວຍຄວາມສະດວກໃນການນຳໃຊ້ຊັບພະຍາກອນ Server ແລະ API ເພື່ອຮັບປະກັນໃຫ້ລະບົບເຮັດວຽກໄດ້ຢ່າງ ")
set_font(r_b1, size=12)
r_b2 = p_body.add_run("ວ່ອງໄວ, ລື່ນໄຫຼ, ເປີດໜ້າເວັບໄດ້ທັນທີ ແລະ ບໍ່ມີອາການຄ້າງ (High Performance & Zero-Downtime)")
set_font(r_b2, bold=True, size=12)
r_b3 = p_body.add_run(" ໃນຊ່ວງທີ່ມີຜູ້ສະໝັກເຂົ້ານຳໃຊ້ພ້ອມກັນເປັນຈຳນວນຫຼາຍ ດັ່ງລາຍລະອຽດລຸ່ມນີ້:")
set_font(r_b3, size=12)

# Point 1
p_p1 = doc.add_paragraph()
p_p1.paragraph_format.space_before = Pt(6)
p_p1.paragraph_format.space_after = Pt(4)
r_p1 = p_p1.add_run("1. ຂໍນຳໃຊ້ Server (Virtual Machine) ແລະ ພື້ນທີ່ຈັດເກັບຂໍ້ມູນ:")
set_font(r_p1, bold=True, size=12)

p_p1_sub = doc.add_paragraph()
p_p1_sub.paragraph_format.space_after = Pt(6)
p_p1_sub.paragraph_format.first_line_indent = Inches(0.3)
r_p1_sub = p_p1_sub.add_run("ເພື່ອຮອງຮັບຂໍ້ມູນຜູ້ສະໝັກວຽກທາງອອນລາຍ ໂດຍສະເລ່ຍປະມານ ")
set_font(r_p1_sub, size=11.5)
r_p1_sub2 = p_p1_sub.add_run("3,000 ຄົນ/ປີ")
set_font(r_p1_sub2, bold=True, size=11.5)
r_p1_sub3 = p_p1_sub.add_run(", ຂໍອະນຸມັດຊັບພະຍາກອນ Server ຕາມສະເປັກຄວາມໄວສູງ ດັ່ງນີ້:")
set_font(r_p1_sub3, size=11.5)

# Spec Table
specs = [
    ("ລາຍການຊັບພະຍາກອນ (Resources)", "ລາຍລະອຽດສະເປັກທີ່ຕ້ອງການ (High Performance Specs)", "ໝາຍເຫດ"),
    ("CPU", "4 - 8 vCPUs (High Frequency)", "ປະມວນຜົນສ້າງໄຟລ໌ PDF ແລະ ແປງຮູບ 3x4 ໄດ້ໄວ"),
    ("RAM", "8 GB - 16 GB", "ຮອງຮັບການເຂົ້າໃຊ້ງານພ້ອມກັນຫຼາຍຮ້ອຍຄົນໂດຍບໍ່ຄ້າງ"),
    ("Hard Disk (Storage)", "200 GB (NVMe SSD)", "ອ່ານ-ຂຽນຂໍ້ມູນໄວສູງສຸດ ສຳລັບຖານຂໍ້ມູນ ແລະ PDF"),
    ("Network Speed", "1 Gbps Port", "ໂຫຼດໜ້າເວັບໄວ, ອັບໂຫຼດ/ດາວໂຫຼດເອກະສານບໍ່ສະດຸດ"),
    ("Operating System", "Linux (Ubuntu 22.04 LTS / Rocky Linux 9)", "ລະບົບປະຕິບັດການທີ່ໄວ, ໝັ້ນຄົງ ແລະ ປອດໄພ"),
    ("Software Runtime", "Node.js (v18+) ແລະ Nginx Web Server", "ແລ່ນ Backend & Frontend ພ້ອມບີບອັດ Cache"),
    ("Database", "PostgreSQL / MySQL (ຫຼື MongoDB Community)", "ໃຊ້ຖານຂໍ້ມູນມາດຕະຖານຂອງ ລທລ")
]

table_spec = doc.add_table(rows=len(specs), cols=3)
table_spec.alignment = WD_TABLE_ALIGNMENT.CENTER
table_spec.autofit = False

col_widths = [Inches(1.8), Inches(2.7), Inches(2.2)]

for row_idx, row_data in enumerate(specs):
    row = table_spec.rows[row_idx]
    for col_idx, text in enumerate(row_data):
        cell = row.cells[col_idx]
        cell.width = col_widths[col_idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        
        # Style Header vs Data
        if row_idx == 0:
            shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E21C25"/>')
            cell._tc.get_or_add_tcPr().append(shading_elm)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            set_font(r, bold=True, size=10, color=RGBColor(255, 255, 255))
        else:
            if row_idx % 2 == 1:
                shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8F9FA"/>')
                cell._tc.get_or_add_tcPr().append(shading_elm)
            r = p.add_run(text)
            if col_idx == 0:
                set_font(r, bold=True, size=10)
            elif col_idx == 1:
                set_font(r, bold=True, size=10, color=RGBColor(190, 20, 30))
            else:
                set_font(r, size=9.5)

# Set table borders
tblPr = table_spec._tbl.tblPr
tblBorders = parse_xml(
    f'<w:tblBorders {nsdecls("w")}>'
    f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>'
    f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>'
    f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/>'
    f'<w:insideV w:val="none"/>'
    f'<w:left w:val="none"/>'
    f'<w:right w:val="none"/>'
    f'</w:tblBorders>'
)
tblPr.append(tblBorders)

# Note below table
p_note = doc.add_paragraph()
p_note.paragraph_format.space_before = Pt(4)
p_note.paragraph_format.space_after = Pt(8)
r_n1 = p_note.add_run("*(ໝາຍເຫດ: ໃນກໍລະນີຜູ້ສະໝັກສຳພາດບໍ່ຜ່ານ ລະບົບຈະລຶບໄຟລ໌ອອກທັນທີ ເພື່ອປະຢັດພື້ນທີ່; ສ່ວນຜູ້ທີ່ຜ່ານການຄັດເລືອກ ຂໍ້ມູນຈະຖືກສົ່ງໄປຈັດເກັບໄວ້ໃນ ")
set_font(r_n1, italic=True, size=10.5)
r_n2 = p_note.add_run("LTC Drive")
set_font(r_n2, bold=True, italic=True, size=10.5)
r_n3 = p_note.add_run(" ຂອງ ລທລ).*")
set_font(r_n3, italic=True, size=10.5)

# Point 2
p_p2 = doc.add_paragraph()
p_p2.paragraph_format.space_before = Pt(4)
p_p2.paragraph_format.space_after = Pt(3)
r_p2 = p_p2.add_run("2. ຂໍເຊື່ອມຕໍ່ລະບົບແຈ້ງເຕືອນ (API & Notifications):")
set_font(r_p2, bold=True, size=12)

p_p2_1 = doc.add_paragraph()
p_p2_1.paragraph_format.space_after = Pt(2)
p_p2_1.paragraph_format.left_indent = Inches(0.3)
r_p2_1a = p_p2_1.add_run("- ຂໍນຳໃຊ້ ")
set_font(r_p2_1a, size=11.5)
r_p2_1b = p_p2_1.add_run("ອີເມວກາງຂອງ ລທລ (ເຊັ່ນ: recruitment@laotel.com)")
set_font(r_p2_1b, bold=True, size=11.5)
r_p2_1c = p_p2_1.add_run(" ພ້ອມເປີດ Port SMTP (Port 587/465) ເພື່ອໃຫ້ລະບົບສົ່ງອີເມວແຈ້ງຜົນການສະໝັກ ແລະ ແຈ້ງມື້ນັດສຳພາດໃຫ້ແກ່ຜູ້ສະໝັກແບບອັດໂນມັດ.")
set_font(r_p2_1c, size=11.5)

p_p2_2 = doc.add_paragraph()
p_p2_2.paragraph_format.space_after = Pt(6)
p_p2_2.paragraph_format.left_indent = Inches(0.3)
r_p2_2a = p_p2_2.add_run("- ຂໍເຊື່ອມຕໍ່ ")
set_font(r_p2_2a, size=11.5)
r_p2_2b = p_p2_2.add_run("API SMS Gateway")
set_font(r_p2_2b, bold=True, size=11.5)
r_p2_2c = p_p2_2.add_run(" ຂອງ ລທລ (ຖ້າມີ) ເພື່ອສົ່ງຂໍ້ຄວາມແຈ້ງເຕືອນມື້ນັດສຳພາດຜ່ານເບີໂທລະສັບມືຖືຂອງຜູ້ສະໝັກ.")
set_font(r_p2_2c, size=11.5)

# Point 3
p_p3 = doc.add_paragraph()
p_p3.paragraph_format.space_before = Pt(4)
p_p3.paragraph_format.space_after = Pt(3)
r_p3 = p_p3.add_run("3. ຂໍຕິດຕັ້ງຊື່ Domain / Subdomain:")
set_font(r_p3, bold=True, size=12)

p_p3_sub = doc.add_paragraph()
p_p3_sub.paragraph_format.space_after = Pt(10)
p_p3_sub.paragraph_format.left_indent = Inches(0.3)
r_p3_a = p_p3_sub.add_run("ຂໍອະນຸໂລມຕັ້ງຊື່ Subdomain ພາຍໃຕ້ໂດເມນບໍລິສັດ ຄື: ")
set_font(r_p3_a, size=11.5)
r_p3_b = p_p3_sub.add_run("LTC-Recruitment.laotel.com")
set_font(r_p3_b, bold=True, size=11.5, color=RGBColor(227, 28, 37))
r_p3_c = p_p3_sub.add_run(" ພ້ອມຕິດຕັ້ງໃບຢັ້ງຢືນຄວາມປອດໄພ SSL (HTTPS) ເພື່ອໃຫ້ຜູ້ສະໝັກທົ່ວໄປສາມາດເຂົ້າເຖິງໜ້າເວັບໄຊໄດ້ຢ່າງສະດວກ, ປອດໄພ ແລະ ມີຄວາມໜ້າເຊື່ອຖືສູງ.")
set_font(r_p3_c, size=11.5)

# Closing
p_close = doc.add_paragraph()
p_close.paragraph_format.space_before = Pt(6)
p_close.paragraph_format.space_after = Pt(8)
p_close.paragraph_format.first_line_indent = Inches(0.5)
r_close = p_close.add_run("ດັ່ງນັ້ນ, ຈຶ່ງຮຽນສະເໜີມາຍັງທ່ານ ເພື່ອພິຈາລະນາ ແລະ ອະນຸມັດຕາມຄວາມເໝາະສົມດ້ວຍ.")
set_font(r_close, size=12)

p_sign_top = doc.add_paragraph()
p_sign_top.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_sign_top.paragraph_format.space_before = Pt(6)
p_sign_top.paragraph_format.space_after = Pt(20)
r_st = p_sign_top.add_run("ຮຽນມາດ້ວຍຄວາມເຄົາລົບ ແລະ ນັບຖືຢ່າງສູງ")
set_font(r_st, size=12)

# Signature Table
table_sign = doc.add_table(rows=1, cols=2)
table_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
table_sign.autofit = False

cell_s1 = table_sign.cell(0, 0)
cell_s1.width = Inches(3.2)
p_s1 = cell_s1.paragraphs[0]
p_s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_s1_t = p_s1.add_run("ຫົວໜ້າພະແນກບໍລິຫານ")
set_font(r_s1_t, bold=True, size=12)

cell_s2 = table_sign.cell(0, 1)
cell_s2.width = Inches(3.2)
p_s2 = cell_s2.paragraphs[0]
p_s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_s2_t = p_s2.add_run("ຫົວໜ້າພະແນກຈັດຕັ້ງ")
set_font(r_s2_t, bold=True, size=12)

# Save Document
out_path = r"C:\Users\Administrator\Desktop\sa_muk_vk_LTC\LTC_Recruitment_Proposal.docx"
doc.save(out_path)
print("File successfully created at: LTC_Recruitment_Proposal.docx")
